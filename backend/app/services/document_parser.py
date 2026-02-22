"""Document parsing service."""
import math
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt", ".xlsx", ".xls"}
MAX_FILE_SIZE_BYTES = 5 * 1024 * 1024  # 5 MB
MAX_PAGE_COUNT = 5


def parse_document(file_path: str) -> dict:
    """
    Parse document and extract text and metadata.
    Returns: { document_text: str, document_metadata: dict }
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    suffix = path.suffix.lower()
    if suffix not in ALLOWED_EXTENSIONS:
        raise ValueError(
            f"Unsupported format: {suffix}. Allowed: {', '.join(ALLOWED_EXTENSIONS)}"
        )

    metadata = {
        "filename": path.name,
        "type": suffix.lstrip("."),
    }

    if suffix == ".pdf":
        return _parse_pdf(path, metadata)
    if suffix in (".docx", ".doc"):
        return _parse_docx(path, metadata)
    if suffix == ".txt":
        return _parse_txt(path, metadata)
    if suffix in (".xlsx", ".xls"):
        return _parse_xlsx(path, metadata)

    raise ValueError(f"Unsupported format: {suffix}")


def _parse_pdf(path: Path, metadata: dict) -> dict:
    """Parse PDF using PyPDF2."""
    from PyPDF2 import PdfReader

    reader = PdfReader(str(path))
    pages = reader.pages
    text = "\n".join(p.extract_text() or "" for p in pages)
    metadata["page_count"] = len(pages)
    return {"document_text": text, "document_metadata": metadata}


def _parse_docx(path: Path, metadata: dict) -> dict:
    """Parse Word document using python-docx."""
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    paragraphs = [p.text for p in doc.paragraphs]
    text = "\n".join(paragraphs)
    word_count = len(text.split())
    metadata["paragraph_count"] = len(paragraphs)
    metadata["page_count"] = max(1, math.ceil(word_count / 300))
    return {"document_text": text, "document_metadata": metadata}


def _parse_txt(path: Path, metadata: dict) -> dict:
    """Parse plain text file."""
    with open(path, encoding="utf-8", errors="replace") as f:
        text = f.read()
    metadata["page_count"] = max(1, math.ceil(len(text) / 1800))
    return {"document_text": text, "document_metadata": metadata}


def _parse_xlsx(path: Path, metadata: dict) -> dict:
    """Parse Excel file using openpyxl."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    sheet_count = len(wb.worksheets)
    sheets_text = []
    for sheet in wb.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            row_str = "\t".join(str(c) if c is not None else "" for c in row)
            if row_str.strip():
                rows.append(row_str)
        if rows:
            sheets_text.append(f"[Sheet: {sheet.title}]\n" + "\n".join(rows))
    wb.close()
    text = "\n\n".join(sheets_text)
    metadata["page_count"] = max(1, sheet_count)
    return {"document_text": text, "document_metadata": metadata}
